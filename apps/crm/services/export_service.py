from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

COLUMNS = [
    ('Nome', lambda lead: lead.customer_name or 'Sem nome'),
    ('Telefone', lambda lead: lead.customer_phone),
    ('Estágio', lambda lead: lead.get_stage_display()),
    ('Urgência', lambda lead: lead.get_urgency_display()),
    ('Score', lambda lead: lead.qualification_score),
    ('Sentimento', lambda lead: lead.get_sentiment_display()),
    ('Labels', lambda lead: ', '.join(label.name for label in lead.labels.all())),
    ('Responsável', lambda lead: lead.assigned_to.get_full_name() or lead.assigned_to.username
        if lead.assigned_to else ''),
    ('Tipo de evento', lambda lead: lead.event_type.name if lead.event_type else ''),
    ('Data do evento', lambda lead: lead.event_date),
    ('Convidados', lambda lead: lead.guest_count),
    ('Orçamento estimado', lambda lead: lead.estimated_budget),
    ('Aguardando atendimento humano', lambda lead: 'Sim' if lead.escalated_at else 'Não'),
    ('Última interação', lambda lead: _naive(lead.last_interaction_at)),
    ('Criado em', lambda lead: _naive(lead.created_at)),
]


def _naive(value):
    """Strips timezone info: Excel's datetime cells can't hold tz-aware values."""
    return value.replace(tzinfo=None) if value else value


class LeadExportService:
    """Builds an .xlsx workbook from a queryset of leads."""

    @staticmethod
    def to_xlsx(leads):
        """Returns the workbook as raw bytes, one row per lead."""
        wb = Workbook()
        sheet = wb.active
        sheet.title = 'Leads'

        headers = [label for label, _ in COLUMNS]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)

        for lead in leads:
            sheet.append([getter(lead) for _, getter in COLUMNS])

        for index, header in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = max(len(header) + 2, 14)

        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
